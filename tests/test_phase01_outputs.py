"""Tests for the loader, exports, expert form and content validity indices."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.data import cvi, exports
from src.data.expert_form import build_workbook, load_items
from src.data.load import check_headers, column_specs, short_names


def test_column_dictionary_is_complete() -> None:
    names = short_names()
    assert len(names) == 53
    assert len(set(names)) == 53
    assert [c["position"] for c in column_specs()] == list(range(53))


def test_check_headers_rejects_wrong_file() -> None:
    specs = column_specs()
    good = [s["header_start"] + " rest of header" for s in specs]
    check_headers(good, specs)
    with pytest.raises(ValueError, match="does not match"):
        check_headers(["Wrong"] + good[1:], specs)
    with pytest.raises(ValueError, match="Expected 53"):
        check_headers(good[:-1], specs)


def test_option_codes_follow_form_order() -> None:
    assert exports.option_codes("gender") == {"Male": 1, "Female": 2, "Prefer not to say": 3}


def test_spss_syntax_reverses_and_sets_levels() -> None:
    likert = ["se1", "se2", "er4"]
    dictionary = exports.spss_dictionary(likert, [])
    syntax = exports.spss_syntax(
        "x.csv", dictionary, {"a": ["se1", "se2"], "b": ["er4"]}, ["er4"], 2
    )
    assert "RECODE er4 (1=5) (2=4) (3=3) (4=2) (5=1) INTO er4_r." in syntax
    assert "gender ol_school_type (NOMINAL)" in syntax
    assert "RELIABILITY /VARIABLES=er4_r" in syntax


def _filled_workbook(tmp_path: Path, name: str, rating: int) -> Path:
    book = build_workbook(load_items())
    sheet = book[cvi.RATINGS_SHEET]
    for row in range(2, sheet.max_row + 1):
        for col in (7, 8, 9):
            sheet.cell(row=row, column=col, value=rating)
    path = tmp_path / name
    book.save(path)
    return path


def test_expert_workbook_round_trip(tmp_path: Path) -> None:
    path = _filled_workbook(tmp_path, "e1.xlsx", 4)
    ratings = cvi.read_ratings(path)
    assert len(ratings) == len(load_items())
    sheet = load_workbook(path)[cvi.RATINGS_SHEET]
    assert sheet.data_validations.dataValidation  # 1-4 drop-down is present


def test_read_ratings_rejects_blank(tmp_path: Path) -> None:
    book = build_workbook(load_items())
    path = tmp_path / "blank.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="missing or invalid"):
        cvi.read_ratings(path)


def test_cvi_values() -> None:
    def frame(values: list[int]) -> pd.DataFrame:
        return pd.DataFrame(
            {"no": [1, 2], "relevance": values, "clarity": values, "simplicity": values}
        )

    # Two experts: item 1 both relevant (I-CVI 1.0), item 2 one of two (0.5).
    items = cvi.item_cvi([frame([4, 2]), frame([3, 4])], [3, 4])
    assert items["relevance_icvi"].tolist() == [1.0, 0.5]
    # modified kappa: Pc = C(2,2)*0.25 = 0.25 -> (1-0.25)/(1-0.25) = 1; Pc = C(2,1)*0.25 = 0.5 -> 0
    assert items["relevance_kappa"].tolist() == pytest.approx([1.0, 0.0])
    scale = cvi.scale_cvi(items)
    assert scale.loc["relevance", "s_cvi_ave"] == pytest.approx(0.75)
    assert scale.loc["relevance", "s_cvi_ua"] == pytest.approx(0.5)


def test_icvi_threshold_depends_on_panel_size() -> None:
    cfg = {"i_cvi_min_up_to_5_experts": 1.0, "i_cvi_min_6_plus_experts": 0.78}
    assert cvi.icvi_threshold(1, cfg) == 1.0
    assert cvi.icvi_threshold(6, cfg) == 0.78
